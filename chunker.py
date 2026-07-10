"""
Ortak belge parçalama (chunking) modülü.
batch_load.py ve Streamlit UI tarafından kullanılır.
Tüm belgeler aynı akıllı pipeline'dan geçer.
"""
import os
import pymupdf4llm
from markitdown import MarkItDown
from pptx import Presentation
from langchain_text_splitters import MarkdownHeaderTextSplitter, RecursiveCharacterTextSplitter
from langchain_core.documents import Document


_HEADERS = [("#", "H1"), ("##", "H2"), ("###", "H3"), ("####", "H4")]


def chunk_pptx(path: str, source: str, audience: dict) -> list[Document]:
    """PowerPoint dosyasını slayt başına bir Document olarak parçalar."""
    prs = Presentation(path)
    docs = []
    for i, slide in enumerate(prs.slides):
        parts = []
        if slide.shapes.title and slide.shapes.title.text:
            parts.append(f"# {slide.shapes.title.text.strip()}")
        for shape in slide.shapes:
            if hasattr(shape, "text_frame") and shape.text_frame:
                parts.append(shape.text.strip())
        content = "\n\n".join(parts).strip()
        if content:
            docs.append(Document(
                page_content=content,
                metadata={
                    "source":   source,
                    "chunk_no": i + 1,
                    "file_type": "pptx",
                    "audience": audience,
                },
            ))
    return docs


def chunk_text(path: str, source: str, chunk_size: int, chunk_overlap: int,
               audience: dict) -> list[Document]:
    """
    PDF / DOCX / XLSX dosyasını markdown'a çevirip başlık hiyerarşisine göre
    parçalar, küçük chunk'ları akıllıca birleştirir ve audience metadata'sını ekler.
    Hata durumunda boş liste döner (exception yutar).
    """
    ext = os.path.splitext(path)[1].lower()

    try:
        # ── 1. Ham metin çıkar ─────────────────────────────────────────────────
        if ext == ".pdf":
            raw = pymupdf4llm.to_markdown(path, write_images=False)
        else:
            md = MarkItDown()
            raw = md.convert(path).text_content

        # ── 2. Başlıklara göre böl ─────────────────────────────────────────────
        md_docs = MarkdownHeaderTextSplitter(
            headers_to_split_on=_HEADERS, strip_headers=True
        ).split_text(raw)

        # ── 3. Akıllı birleştirme (yetim / küçük chunk'lar) ───────────────────
        merged: list[Document] = []
        temp: Document | None = None

        for doc in md_docs:
            if not doc.page_content.strip():
                continue

            ctx = " > ".join(doc.metadata[h] for _, h in _HEADERS if doc.metadata.get(h))
            if ctx:
                doc.page_content = f"**BAĞLAM:** {ctx}\n\n{doc.page_content}"

            if temp:
                if len(doc.page_content) < 100 and "|" not in doc.page_content:
                    merged.append(temp)
                    temp = doc
                else:
                    doc.page_content = f"{temp.page_content}\n\n{doc.page_content}"
                    merged.append(doc)
                    temp = None
            else:
                if len(doc.page_content) < 250 and "|" not in doc.page_content:
                    temp = doc
                else:
                    merged.append(doc)

        if temp:
            merged.append(temp)

        # ── 4. Büyük chunk'ları recursive split ───────────────────────────────
        rec = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            separators=["\n\n", "\n", " ", ""],
        )
        final: list[Document] = []
        for doc in merged:
            doc.metadata.update({
                "source":    source,
                "file_type": ext.lstrip("."),
                "audience":  audience,
            })
            final.extend(rec.split_documents([doc]))

        return final

    except Exception as e:
        print(f"[chunker] {source}: {e}")
        return []